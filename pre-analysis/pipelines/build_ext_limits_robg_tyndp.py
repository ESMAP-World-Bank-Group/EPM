# -*- coding: utf-8 -*-
"""TYNDP NTC sensitivity for the RoBg family: external limits that grow to the TYNDP 2024 values.

Writes, next to the central file and without touching it:

    epm/input/data_blacksea/trade/pExtTransferLimit_robg_tyndp.csv

and the extracted TYNDP capacities to pre-analysis/output_prices/robg/tyndp_ntc_2024.csv.

Method
------
* Start from trade/pExtTransferLimit_robg.csv (central RoBg file, written by
  build_ext_limits_robg.py). Every row except the five RO/BG borders is copied
  unchanged.
* TYNDP 2024 capacity per border and direction = maximum hourly cross-border
  exchange in the National Trends market simulation (MMStandardOutputFile
  NT2030 and NT2040, Plexos, climate year 2009, sheet "Crossborder exchanges").
  The market model saturates every border at least once a year, so the hourly
  maximum is the NTC the simulation used.
* Capacity per border, direction and year:
    - up to 2025: central value (2025 p99 of the ENTSO-E physical flow), so the
      calibration year is identical to the central run;
    - 2030: max(central, TYNDP NT2030); 2040: max(central, TYNDP NT2040);
    - linear between 2025, 2030 and 2040, flat after 2040, same for all seasons.
  The max() keeps the sensitivity a pure relaxation: Bulgaria-Serbia carried
  744 MW in 2025 while the NT2030 value is 380 MW.
* Trakia-Bulgaria and Trakia-Greece are not touched (kept consistent with the
  external family).

Caveat: the external zones are price takers at the Hungarian price. Larger
limits therefore read as an upper bound on export volumes and corridor value,
not as a forecast.

Usage (from EPM/):  python pre-analysis/pipelines/build_ext_limits_robg_tyndp.py
"""
from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd

_HERE = Path(__file__).resolve()
_PRE_ANALYSIS = _HERE.parents[1]
_EPM = _PRE_ANALYSIS.parent

TRADE_DIR = _EPM / "epm" / "input" / "data_blacksea" / "trade"
CENTRAL = TRADE_DIR / "pExtTransferLimit_robg.csv"
OUT_LIMIT = TRADE_DIR / "pExtTransferLimit_robg_tyndp.csv"
OUT_DIR = _PRE_ANALYSIS / "output_prices" / "robg"
TYNDP_DIR = _PRE_ANALYSIS.parents[1] / "Data" / "TYNDP" / "2024"
TYNDP_FILES = {
    2030: "MMStandardOutputFile_NT2030_Plexos_CY2009_2.5_v40.xlsx",
    2040: "MMStandardOutputFile_NT2040_Plexos_CY2009_2.5_v40.xlsx",
}
SHEET = "Crossborder exchanges"
HEADER_ROW = 11
BASE_YEAR = 2025

# (internal zone, external zone, TYNDP node inside, TYNDP node outside)
BORDERS = [
    ("Romania", "Hungary", "RO00", "HU00"),
    ("Romania", "Serbia", "RO00", "RS00"),
    ("Bulgaria", "Serbia", "BG00", "RS00"),
    ("Bulgaria", "NorthMacedonia", "BG00", "MK00"),
    ("Bulgaria", "Greece", "BG00", "GR00"),
]


def _fail(msg: str) -> None:
    raise SystemExit(f"[build_ext_limits_robg_tyndp] FAIL: {msg}")


def tyndp_capacity(year: int) -> dict:
    """{(zone, zext): (export_mw, import_mw)} from the hourly maxima of one TYNDP file."""
    from openpyxl import load_workbook

    path = TYNDP_DIR / TYNDP_FILES[year]
    if not path.exists():
        _fail(f"{path} not found")
    ws = load_workbook(path, read_only=True)[SHEET]
    rows = ws.iter_rows(min_row=HEADER_ROW, values_only=True)
    head = list(next(rows))
    data = [r for r in rows if r[0] is not None]
    out = {}
    for z, x, inside, outside in BORDERS:
        if f"{inside}->{outside}" in head:
            col, sign = head.index(f"{inside}->{outside}"), 1.0
        elif f"{outside}->{inside}" in head:
            col, sign = head.index(f"{outside}->{inside}"), -1.0
        else:
            _fail(f"no column for {inside}-{outside} in {path.name}")
        s = sign * pd.to_numeric(pd.Series([r[col] for r in data]), errors="coerce").fillna(0.0)
        if len(s) < 8700:
            _fail(f"{inside}-{outside}: {len(s)} hours in {path.name}")
        out[(z, x)] = (float(s.max()), float(-s.min()))
    return out


def main() -> int:
    tyndp = {y: tyndp_capacity(y) for y in TYNDP_FILES}

    raw = CENTRAL.read_text(encoding="utf-8").splitlines()
    header = raw[0].split(",")
    years = [int(y) for y in header[4:]]
    targets = {(z, x) for z, x, _, _ in BORDERS}

    central = {}
    for ln in raw[1:]:
        p = ln.split(",")
        if (p[0], p[1]) in targets:
            central.setdefault((p[0], p[1], p[3]), set()).update(float(v) for v in p[4:])
    for key, vals in central.items():
        if len(vals) != 1:
            _fail(f"central limit not flat for {key}: {sorted(vals)}")
    central = {k: v.pop() for k, v in central.items()}

    table, path = [], {}
    for z, x, _, _ in BORDERS:
        for d, i in (("Export", 0), ("Import", 1)):
            c = central[(z, x, d)]
            t30, t40 = max(c, tyndp[2030][(z, x)][i]), max(c, tyndp[2040][(z, x)][i])
            path[(z, x, d)] = np.round(np.interp(years, [BASE_YEAR, 2030, 2040], [c, t30, t40])).astype(int)
            table.append({"zone": z, "zext": x, "direction": d, "central_mw": c,
                          "tyndp_2030_mw": tyndp[2030][(z, x)][i], "tyndp_2040_mw": tyndp[2040][(z, x)][i],
                          "used_2030_mw": t30, "used_2040_mw": t40})

    lines = [raw[0]]
    for ln in raw[1:]:
        p = ln.split(",")
        if (p[0], p[1]) in targets:
            lines.append(",".join(p[:4] + [str(v) for v in path[(p[0], p[1], p[3])]]))
        else:
            lines.append(ln)
    # Same LF convention as the central file (write_bytes: no newline translation on Windows).
    OUT_LIMIT.write_bytes(("\n".join(lines) + "\n").encode("utf-8"))

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    table = pd.DataFrame(table)
    table.to_csv(OUT_DIR / "tyndp_ntc_2024.csv", index=False)
    with pd.option_context("display.width", 200, "display.max_columns", 20):
        print(table.to_string(index=False))
    print(f"\n  wrote {OUT_LIMIT.name}: {len(lines) - 1} rows")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
