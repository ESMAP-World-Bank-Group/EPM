# -*- coding: utf-8 -*-
"""Trade prices for the RoBg scenarios, where Romania and Bulgaria are internal.

Writes, next to the reference files and without touching them:

    epm/input/data_blacksea/trade/pTradePrice_robg.csv
    epm/input/data_blacksea/trade/pTradePriceExport_robg.csv

and the intermediate artefacts to pre-analysis/output_prices/robg/.

What changes against pTradePrice_eu_central.csv
-----------------------------------------------
* Romania and Bulgaria rows are dropped: both become internal zones.
* Six external zones around Romania and Bulgaria are priced by one chain:
  Hungary, Serbia, NorthMacedonia, Greece, Moldova and Ukraine.

      P(z,q,d,t,y) = L_HU(y) x r(z,y) x S(z,q,d,t) x T(z,y,q,t) / D(z,q,t)

    L_HU  eu_price_level.py, observed 2024 joined to TYNDP 2024 National Trends
          (HU00), CENTRAL trajectory. The level of every zone hangs on it.
    r     level of the zone relative to Hungary. Model year 2025: observed 2024
          day-ahead mean GR / HU for Greece, RO / HU for Moldova, 1 for the
          zones without a cached day-ahead series (Serbia, North Macedonia,
          Ukraine).
          2030 and 2040: TYNDP NT annual mean marginal cost of the node over
          HU00. Linear in between, flat after 2040.
    S     eu_price.py: 2023 day-ahead prices (ENTSO-E) sampled by calendar date
          on the representative days, weighted mean 1. Greek series for
          Greece, Hungarian series for the five others (Moldova included: the
          2023 Romanian days miss gate G2bis, P10 19.2 % of the mean against
          15 %, as in the reference chain).
    D     pHours-weighted mean of S within each (quarter, hour).
    T     target quarter x hour profile, weighted mean 1. Model year 2025: the
          observed 2024 day-ahead profile (GR for Greece, RO for Moldova, HU
          otherwise).
          2030 and 2040: the TYNDP NT hourly marginal cost of the node. Linear
          in between, flat after 2040.

  S x T / D keeps the day-to-day spread of the observed 2023 days and moves the
  average daily shape onto the target. The weighted mean of T / D x S over the
  year is exactly 1, so the annual mean stays L_HU x r (gate below).
  Why the profile moves: 2023 prices carry the 2023 fleet. By 2040 the TYNDP
  fleet has much more solar in the region, so the midday hours fall well
  below the evening ones. A 2023 shape held to 2050 understates that dip.
  Why Romania stands in for Moldova (since 2026-10-20): Moldova imports most
  of its electricity from Romania and has no ENTSO-E day-ahead series, so its
  2025 price is the Romanian one. The 2024 Romanian mean sits 2.7 percent
  above the Hungarian one, with an hourly correlation of 0.94. Level r and
  profile T are Romanian; only the day-to-day spread S stays Hungarian.
  Ukraine keeps the Hungarian stand-in: no Ukrainian series is cached (caveat).
* Netback, the same on every border: the seller bears an AC line loss of
  2.5 % (bareme of 2026-08-19, `lambda` in config/eu_price_netback.csv).
  W = 0: the ITC perimeter fee falls on non-ITC parties. Romania, Bulgaria,
  Hungary, Greece, Serbia and North Macedonia sit inside the ITC mechanism,
  Ukraine since July 2024; Moldova is treated the same way, the fee being
  second order (1.10 EUR/MWh in 2026). C = 0, as in the reference files: the
  CBAM levy is a separate scenario family.
* Every other zone (Iran, Iraq, Syria, Russia, Kazakhstan) is copied byte for
  byte from the reference files.

TYNDP clock: the hourly sheet runs on UTC, as the model does (solar output
peaks at hour 10 in HU00, RO00 and GR00). TYNDP hour h maps to model t{h+1}.
The hourly marginal costs are cached in output_prices/robg/tyndp_nt_mc_profile.csv,
and each node's annual mean is checked against the 'Avg MC' row of its workbook.

Usage (from EPM/):  python pre-analysis/pipelines/build_trade_prices_robg.py
"""
from __future__ import annotations

import sys
from pathlib import Path

import numpy as np
import pandas as pd

_HERE = Path(__file__).resolve()
_PRE_ANALYSIS = _HERE.parents[1]
_EPM = _PRE_ANALYSIS.parent
sys.path.insert(0, str(_HERE.parent))

import eu_price  # noqa: E402
import eu_price_level  # noqa: E402

HOURS = [f"t{h:02d}" for h in range(1, 25)]
KEY = ["zext", "q", "d", "year"]

TRADE_DIR = _EPM / "epm" / "input" / "data_blacksea" / "trade"
REF_IMPORT = TRADE_DIR / "pTradePrice_eu_central.csv"
REF_EXPORT = TRADE_DIR / "pTradePriceExport_eu_central.csv"
OUT_IMPORT = TRADE_DIR / "pTradePrice_robg.csv"
OUT_EXPORT = TRADE_DIR / "pTradePriceExport_robg.csv"
ZEXT_ROBG = TRADE_DIR / "zext_robg.csv"
OUT_DIR = _PRE_ANALYSIS / "output_prices" / "robg"
TYNDP_CACHE = OUT_DIR / "tyndp_nt_mc_profile.csv"

CACHE_DIR = _PRE_ANALYSIS.parents[1] / "Data" / "cache_entso_e"
TYNDP_DIR = _PRE_ANALYSIS.parents[1] / "Data" / "TYNDP" / "2024"
REPDAYS_DIR = _PRE_ANALYSIS / "representative_days" / "output" / "blacksea"
DEFLATORS = _PRE_ANALYSIS / "config" / "price_deflators.csv"
NETBACK_CFG = _PRE_ANALYSIS / "config" / "eu_price_netback.csv"

DROPPED = ("Romania", "Bulgaria")
LEVEL_ZONE = "Hungary"
# zone -> (TYNDP node, ENTSO-E bidding zone whose day-ahead series stands in)
PRICED = {
    "Hungary": ("HU00", "HU"),
    "Serbia": ("RS00", "HU"),
    "NorthMacedonia": ("MK00", "HU"),
    "Greece": ("GR00", "GR"),
    "Moldova": ("MD00", "RO"),
    "Ukraine": ("UA00", "HU"),
}
# Day-ahead series whose 2023 days fail gate G2bis: S is taken from the stand-in.
SHAPE_STAND_IN = {"RO": "HU"}
SHAPE_YEAR = 2023         # observed days behind S
OBS_YEAR = 2024           # observed profile and level ratio, anchored on model year 2025
ANCHOR_YEAR = 2025
TYNDP_YEARS = (2030, 2040)
LEVEL = "CENTRAL"
FLOOR_USD = 0.01          # same strictly positive floor as eu_price_netback.py
DECIMALS = 3
D_MIN = 0.2               # a (quarter, hour) cell this cheap in S would blow up T / D
MONTHS = ["JAN", "FEB", "MAR", "APR", "MAY", "JUN",
          "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"]


def _fail(msg: str) -> None:
    raise SystemExit(f"[build_trade_prices_robg] FAIL: {msg}")


# ── Reference pipelines: S for HU and GR, L for HU ─────────────────────────
def build_shape_level() -> tuple[pd.DataFrame, pd.DataFrame]:
    codes = sorted({SHAPE_STAND_IN.get(c, c) for _, c in PRICED.values()})
    names = {"HU": "Hungary", "GR": "Greece"}
    eu_price.ZONES = {c: names[c] for c in codes}
    eu_price_level.ZONES = {"HU": LEVEL_ZONE}
    eu_price_level.TYNDP_CODES = {LEVEL_ZONE: "HU00"}

    qc_s = eu_price.run(CACHE_DIR, REPDAYS_DIR, OUT_DIR, SHAPE_YEAR, 15.0, "Georgia")
    if qc_s["G2bis_failed_zones"]:
        _fail(f"shape gate G2bis failed: {qc_s['G2bis_failed_zones']}")
    qc_l = eu_price_level.run(CACHE_DIR, TYNDP_DIR, DEFLATORS, OUT_DIR,
                              2022, [2024], [2023, 2024])
    # G3 (2023 back-cast) and G3ter (ladder order) are diagnostics in the
    # reference chain too: the promoted eu_central files carry G3 misses for
    # Bulgaria and Greece and a G3ter inversion for Romania. Reported, not fatal.
    if qc_l["failed"]:
        print(f"  !! level diagnostics not met (same status as the reference "
              f"zones, recorded in qc_level_L.json): {qc_l['failed']}")

    shape = pd.read_csv(OUT_DIR / "shape_S.csv")
    level = pd.read_csv(OUT_DIR / "level_L.csv")
    level = level[(level["zone"] == LEVEL_ZONE) & (level["scenario"] == LEVEL)]
    if level.empty:
        _fail(f"no {LEVEL} level for {LEVEL_ZONE}")
    return shape, level[["year", "L_eur2024"]].reset_index(drop=True)


# ── TYNDP NT hourly marginal cost ───────────────────────────────────────────
def _tyndp_book(year: int) -> Path:
    hits = [p for p in sorted(TYNDP_DIR.glob(f"MMStandardOutputFile_NT{year}_*.xls*"))
            if not any(t in p.stem.lower().split("_")
                       for t in eu_price_level.OTHER_COMMODITIES)]
    if len(hits) != 1:
        _fail(f"NT{year}: expected one electricity workbook, found {[p.name for p in hits]}")
    return hits[0]


def _read_tyndp_hourly(year: int, nodes: list[str]) -> pd.DataFrame:
    """Hourly marginal cost per node, long (year, node, q, t, mc), plus the
    'Avg MC' row of the workbook for the check against its own summary."""
    from openpyxl import load_workbook

    path = _tyndp_book(year)
    print(f"  reading {path.name} (hourly sheet, a few minutes)")
    wb = load_workbook(path, read_only=True, data_only=True)
    it = wb["Hourly Market Data"].iter_rows(min_row=1, values_only=True)
    head = [next(it) for _ in range(13)]
    avg_row, code_row = head[9], head[12]
    cols = {}
    for n in nodes:
        label = f"{n}_Mgl Cost"
        if label not in code_row:
            _fail(f"{label} absent from {path.name}")
        cols[n] = code_row.index(label)

    recs = []
    for r in it:
        code = r[1]
        if not isinstance(code, str) or code[2:5] not in MONTHS:
            continue
        recs.append([code[2:5], int(code[5:7])] + [r[c] for c in cols.values()])
    wb.close()
    df = pd.DataFrame(recs, columns=["mon", "h"] + nodes)
    if len(df) < 8736:
        _fail(f"{path.name}: {len(df)} hourly rows, expected at least 8736")
    df["q"] = "Q" + ((df["mon"].map(MONTHS.index) // 3) + 1).astype(str)
    df["t"] = df["h"].map(lambda h: f"t{h + 1:02d}")
    long = df.melt(id_vars=["q", "t"], value_vars=nodes, var_name="node", value_name="mc")
    long["mc"] = pd.to_numeric(long["mc"], errors="coerce")
    if long["mc"].isna().any():
        _fail(f"{path.name}: missing hourly marginal costs")
    # The workbook's own average excludes the 3 000 EUR/MWh scarcity value.
    for n in nodes:
        mine = long.loc[(long["node"] == n) & (long["mc"] < 3000), "mc"].mean()
        theirs = float(avg_row[cols[n]])
        if abs(mine - theirs) > 0.05:
            _fail(f"{n} NT{year}: hourly mean {mine:.3f} vs Avg MC row {theirs:.3f}")
    prof = (long[long["mc"] < 3000].groupby(["node", "q", "t"])["mc"]
            .agg(["mean", "size"]).reset_index()
            .rename(columns={"mean": "mc", "size": "n_hours"}))
    prof.insert(0, "year", year)
    return prof


def tyndp_profiles() -> pd.DataFrame:
    nodes = [n for n, _ in PRICED.values()]
    if TYNDP_CACHE.exists():
        cached = pd.read_csv(TYNDP_CACHE)
        have = {(int(y), n) for y, n in zip(cached["year"], cached["node"])}
        if have == {(y, n) for y in TYNDP_YEARS for n in nodes}:
            print(f"  TYNDP hourly marginal cost from cache {TYNDP_CACHE.name}")
            return cached
    prof = pd.concat([_read_tyndp_hourly(y, nodes) for y in TYNDP_YEARS], ignore_index=True)
    prof.to_csv(TYNDP_CACHE, index=False, float_format="%.4f")
    print(f"  cached {TYNDP_CACHE.name}")
    return prof


# ── Observed 2024 profile and level ─────────────────────────────────────────
def observed(code: str) -> tuple[pd.DataFrame, float]:
    s = eu_price.load_prices(CACHE_DIR, code, OBS_YEAR)
    df = pd.DataFrame({"p": s.to_numpy()})
    df["q"] = "Q" + (((s.index.month - 1) // 3) + 1).astype(str)
    df["t"] = [f"t{h + 1:02d}" for h in s.index.hour]
    prof = df.groupby(["q", "t"])["p"].mean().rename("v").reset_index()
    return prof, float(s.mean())


# ── Profile and ratio paths ─────────────────────────────────────────────────
def _normalise(prof: pd.DataFrame, wqt: pd.DataFrame) -> pd.Series:
    m = prof.merge(wqt, on=["q", "t"], how="right")
    if m["v"].isna().any():
        _fail("a target profile misses a (quarter, hour) cell")
    v = m.set_index(["q", "t"])["v"]
    return v / ((v * m.set_index(["q", "t"])["W"]).sum() / m["W"].sum())


def _blend(points: dict[int, object], year: int):
    """Linear between anchor years, flat outside them."""
    ys = sorted(points)
    if year <= ys[0]:
        return points[ys[0]]
    if year >= ys[-1]:
        return points[ys[-1]]
    if year in points:
        return points[year]
    hi = next(k for k in ys if k >= year)
    lo = max(k for k in ys if k <= year)
    f = (year - lo) / (hi - lo)
    return points[lo] * (1 - f) + points[hi] * f


def price_zone(zone: str, shape: pd.DataFrame, level: pd.DataFrame,
               tyndp: pd.DataFrame, obs: dict, w: pd.DataFrame,
               wqt: pd.DataFrame, rate: float) -> tuple[pd.DataFrame, list, pd.DataFrame]:
    node, code = PRICED[zone]
    s_zone = {"HU": "Hungary", "GR": "Greece"}[SHAPE_STAND_IN.get(code, code)]
    s = (shape[shape["zone"] == s_zone]
         .melt(id_vars=["q", "d"], value_vars=HOURS, var_name="t", value_name="s")
         .merge(w, on=["q", "d", "t"]))
    d_base = (s.assign(sw=s["s"] * s["w"]).groupby(["q", "t"])[["sw", "w"]].sum()
              .pipe(lambda g: g["sw"] / g["w"]))
    if d_base.min() < D_MIN:
        _fail(f"{zone}: D(q,t) falls to {d_base.min():.3f} at {d_base.idxmin()}")

    targets = {ANCHOR_YEAR: _normalise(obs[code][0], wqt)}
    hu_mean = {y: tyndp[(tyndp["year"] == y) & (tyndp["node"] == "HU00")]
               .pipe(lambda g: (g["mc"] * g["n_hours"]).sum() / g["n_hours"].sum())
               for y in TYNDP_YEARS}
    ratios = {ANCHOR_YEAR: obs[code][1] / obs["HU"][1]}
    for y in TYNDP_YEARS:
        t = tyndp[(tyndp["year"] == y) & (tyndp["node"] == node)]
        targets[y] = _normalise(t.rename(columns={"mc": "v"})[["q", "t", "v"]], wqt)
        ratios[y] = ((t["mc"] * t["n_hours"]).sum() / t["n_hours"].sum()) / hu_mean[y]

    rows, trace = [], []
    keys = s.set_index(["q", "t"]).index
    for _, lv in level.iterrows():
        y = int(lv["year"])
        tgt = _blend(targets, y)
        r = _blend(ratios, y)
        factor = (tgt / d_base).reindex(keys).to_numpy()
        unit = s["s"].to_numpy() * factor
        mean_unit = float((unit * s["w"]).sum() / s["w"].sum())
        if abs(mean_unit - 1.0) > 1e-6:
            _fail(f"{zone} {y}: weighted mean of S x T / D is {mean_unit:.9f}")
        rows.append(s[["q", "d", "t"]].assign(year=y, p=unit * lv["L_eur2024"] * r * rate))
        trace.append({"zext": zone, "year": y, "ratio_to_HU": round(r, 4),
                      "L_HU_eur2024": round(float(lv["L_eur2024"]), 3),
                      "mean_usd": round(float(lv["L_eur2024"]) * r * rate, 3)})
    long = pd.concat(rows, ignore_index=True)
    wide = long.pivot_table(index=["q", "d", "year"], columns="t", values="p").reset_index()
    wide.columns.name = None
    tgt_out = pd.concat([v.rename("T").reset_index().assign(zext=zone, anchor=y)
                         for y, v in targets.items()], ignore_index=True)
    return wide[["q", "d", "year"] + HOURS], trace, tgt_out


def netback(hub: pd.DataFrame, lam: float) -> pd.DataFrame:
    exp = hub.copy()
    exp[HOURS] = np.maximum(hub[HOURS].to_numpy() * (1.0 - lam), FLOOR_USD)
    return exp


def assemble(ref_path: Path, priced: dict[str, pd.DataFrame]) -> pd.DataFrame:
    """Reference rows minus Romania, Bulgaria and the priced zones, plus the
    priced blocks in the (q, d, year) order of the reference Greece block."""
    ref = pd.read_csv(ref_path, dtype=str)
    if list(ref.columns) != KEY + HOURS:
        _fail(f"{ref_path.name}: unexpected header")
    order = ref[ref["zext"] == "Greece"][["q", "d", "year"]].copy()
    order["year"] = order["year"].astype(int)
    kept = ref[~ref["zext"].isin(DROPPED + tuple(priced))]

    blocks = [kept]
    for z, frame in priced.items():
        b = order.merge(frame, on=["q", "d", "year"], how="left")
        if b[HOURS].isna().any().any():
            _fail(f"{z}: {int(b[HOURS].isna().any(axis=1).sum())} keys without a price")
        b.insert(0, "zext", z)
        b[HOURS] = b[HOURS].map(lambda v: f"{v:.{DECIMALS}f}")
        b["year"] = b["year"].astype(str)
        blocks.append(b[KEY + HOURS])
    return pd.concat(blocks, ignore_index=True)


def write_like(df: pd.DataFrame, ref_path: Path, out_path: Path) -> None:
    """Same line ending and BOM convention as the reference file."""
    raw = ref_path.read_bytes()
    eol = "\r\n" if b"\r\n" in raw[:4096] else "\n"
    enc = "utf-8-sig" if raw[:3] == b"\xef\xbb\xbf" else "utf-8"
    df.to_csv(out_path, index=False, lineterminator=eol, encoding=enc)


def main() -> int:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    shape, level = build_shape_level()
    tyndp = tyndp_profiles()
    obs = {c: observed(c) for c in sorted({c for _, c in PRICED.values()})}

    w = eu_price.hour_weights(pd.read_csv(REPDAYS_DIR / "pHours.csv"))
    wqt = w.groupby(["q", "t"])["w"].sum().rename("W").reset_index()
    cfg = pd.read_csv(NETBACK_CFG, comment="#")
    lam = float(cfg[(cfg["series"] == "lambda") & (cfg["year"] == 0)]["value"].iloc[0])
    dfl = pd.read_csv(DEFLATORS, comment="#")
    rate = float(dfl[(dfl["series"] == "eur_usd") & (dfl["year"] == 2024)]["value"].iloc[0])

    imp, exp, trace, tgts = {}, {}, [], []
    for zone in PRICED:
        hub, tr, tg = price_zone(zone, shape, level, tyndp, obs, w, wqt, rate)
        imp[zone], exp[zone] = hub, netback(hub, lam)
        trace += tr
        tgts.append(tg)

    def _long(d: dict) -> pd.DataFrame:
        return pd.concat([f.assign(zext=z) for z, f in d.items()], ignore_index=True)[
            ["zext", "q", "d", "year"] + HOURS]

    _long(imp).to_csv(OUT_DIR / "ext_import_usd.csv", index=False, float_format="%.4f")
    _long(exp).to_csv(OUT_DIR / "ext_export_usd.csv", index=False, float_format="%.4f")
    pd.DataFrame(trace).to_csv(OUT_DIR / "ext_level_ratio.csv", index=False)
    pd.concat(tgts, ignore_index=True)[["zext", "anchor", "q", "t", "T"]].to_csv(
        OUT_DIR / "ext_target_profile.csv", index=False, float_format="%.5f")

    a = assemble(REF_IMPORT, imp)
    b = assemble(REF_EXPORT, exp)
    zext = set(pd.read_csv(ZEXT_ROBG).iloc[:, 0].astype(str))
    for p, df in ((OUT_IMPORT, a), (OUT_EXPORT, b)):
        missing = zext - set(df["zext"])
        if missing:
            _fail(f"{p.name}: external zones without a price {sorted(missing)}")
    write_like(a, REF_IMPORT, OUT_IMPORT)
    write_like(b, REF_EXPORT, OUT_EXPORT)

    print(f"\n  lambda {lam:.3f} | EUR->USD {rate:.4f}")
    tr = pd.DataFrame(trace)
    print(tr[tr["year"].isin([2025, 2030, 2040])]
          .pivot(index="zext", columns="year", values="mean_usd").round(1).to_string())
    # Shape check: Q2 midday (t11 to t14 UTC) over Q4 evening (t16 to t19).
    for zone in ("Hungary", "Greece", "Ukraine"):
        day_w = w[w["t"] == "t01"][["q", "d", "w"]]
        h = imp[zone].merge(day_w, on=["q", "d"])
        out = []
        for y in (2025, 2030, 2040):
            g = h[h["year"] == y]

            def wm(q, cols):
                x = g[g["q"] == q]
                return float((x[cols].mean(axis=1) * x["w"]).sum() / x["w"].sum())
            out.append(f"{y} Q2 midday {wm('Q2', ['t11', 't12', 't13', 't14']):.0f} "
                       f"/ Q4 evening {wm('Q4', ['t16', 't17', 't18', 't19']):.0f}")
        print(f"  {zone:14s} " + " | ".join(out))
    for p, df in ((OUT_IMPORT, a), (OUT_EXPORT, b)):
        print(f"  wrote {p.name}: {len(df)} rows, zones {sorted(df['zext'].unique())}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
