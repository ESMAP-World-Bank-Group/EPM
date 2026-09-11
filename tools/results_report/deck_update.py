# -*- coding: utf-8 -*-
"""Point every regenerated chart in the results deck at the current run.

Only the image relationship of each picture is rewritten, so position, crop and
z-order survive untouched.  Pictures the generator cannot produce are left in
place and flagged on the slide instead of being silently carried over: a stale
chart that looks current is worse than an obviously missing one.

    python deck_update.py                        # writes the September deck
    python deck_update.py --deck other.pptx --out copy.pptx
    python deck_update.py --dry-run              # report, change nothing
"""
import argparse
import shutil
import struct
from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

import runcfg

ROOT = runcfg.ROOT.parent                              # .../blacksea_2026
DECK = ROOT / "BlackSea_regional_power_trade_followup_September2026_results.pptx"
PNGS = ROOT / "Data" / "results" / "slides"

PICTURE = 13                                           # MSO_SHAPE_TYPE.PICTURE
FLAG = "stale-marker"                                  # our own overlays
ASPECT_TOL = 0.02                                      # 2 %, PowerPoint crops

# slide (1-based) -> shape name -> chart file.  Every pair below was matched
# against the deck by md5 of the embedded image, except slide 11 Picture 13
# whose stale copy was identified by image correlation (0.96 against
# region_generation, 0.26 against the nearest other candidate).
CHARTS = {
    3:  {"Picture 20": "dispatch_georgia_2025_2035.png"},
    4:  {"Picture 22": "trade_georgia.png",
         "Picture 34": "seasonal_georgia.png"},
    5:  {"Picture 4":  "dispatch_azerbaijan_2025_2035.png"},
    6:  {"Picture 22": "trade_azerbaijan.png",
         "Picture 10": "seasonal_azerbaijan.png"},
    7:  {"Picture 6":  "dispatch_armenia_2025_2035.png"},
    8:  {"Picture 22": "trade_armenia.png",
         "Picture 3":  "seasonal_armenia.png"},
    9:  {"Picture 4":  "dispatch_turkiye_2025_2035.png"},
    10: {"Picture 22": "trade_turkiye.png",
         "Picture 5":  "seasonal_turkiye.png"},
    11: {"Picture 7":  "region_maps.png",
         "Picture 13": "region_generation.png"},
    13: {"Picture 10": "bssc_volume_sensitivity.png"},
    14: {"Picture 7":  "bssc_mix_delta.png"},
}

# Pictures no generator can redraw: they are screenshots of the HTML results
# report.  Rebuild it with build.py and recapture, or write the matplotlib
# equivalents, then move the entry into CHARTS above.
STALE = {
    3:  {"Picture 2":  "capacity vs NDP"},
    4:  {"Picture 28": "flow maps"},
    5:  {"Picture 2":  "capacity vs NDP"},
    6:  {"Picture 8":  "flow maps"},
    7:  {"Picture 4":  "capacity vs NDP"},
    8:  {"Picture 5":  "flow maps"},
    9:  {"Picture 2":  "capacity vs NDP"},
    10: {"Picture 3":  "flow maps"},
    12: {"Picture 3":  "regional panel", "Picture 5": "regional panel",
         "Picture 8":  "regional panel"},
    13: {"Picture 4":  "BSSC map"},
    14: {"Picture 4":  "BSSC country panel"},
    15: {"Picture 4":  "Free Expansion map"},
}


def png_size(path):
    b = path.read_bytes()[:24]
    return struct.unpack(">II", b[16:24]) if b[:8] == b"\x89PNG\r\n\x1a\n" \
        else (0, 0)


def blob_size(pic):
    b = pic.image.blob[:24]
    return struct.unpack(">II", b[16:24]) if b[:8] == b"\x89PNG\r\n\x1a\n" \
        else (0, 0)


def find(slide, name):
    hits = [sh for sh in slide.shapes
            if sh.shape_type == PICTURE and sh.name == name]
    if len(hits) != 1:
        return None
    return hits[0]


def clear_flags(slide):
    """Drop the markers a previous pass left, so reruns do not stack them."""
    for sh in list(slide.shapes):
        if sh.name == FLAG:
            sh._element.getparent().remove(sh._element)


def flag(slide, pic, why, run):
    """Small red banner over a picture the run could not refresh."""
    w = min(pic.width, Emu(int(2.30 * 914400)))
    box = slide.shapes.add_textbox(pic.left, pic.top, w,
                                   Emu(int(0.17 * 914400)))
    box.name = FLAG
    tf = box.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = Emu(int(0.03 * 914400))
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = "NOT UPDATED  %s  (%s)" % (why, run.replace("simulations_run_", ""))
    r.font.size = Pt(7)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0xD9, 0x53, 0x4F)
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(0xFD, 0xF6, 0xF5)
    box.line.color.rgb = RGBColor(0xD9, 0x53, 0x4F)
    box.line.width = Pt(0.75)


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--deck", default=str(DECK))
    p.add_argument("--out", default=None)
    p.add_argument("--pngs", default=str(PNGS))
    p.add_argument("--dry-run", action="store_true")
    a = p.parse_args()

    deck, pngs = Path(a.deck), Path(a.pngs)
    out = Path(a.out) if a.out else deck
    prs = Presentation(str(deck))
    run = runcfg.DEFAULT_RUN
    done = skipped = flagged = 0

    for idx in sorted(set(CHARTS) | set(STALE)):
        slide = prs.slides[idx - 1]
        clear_flags(slide)
        for name, fn in CHARTS.get(idx, {}).items():
            pic, src = find(slide, name), pngs / fn
            if pic is None:
                print("s%-2d %-11s SKIP  shape not found" % (idx, name))
                skipped += 1
                continue
            if not src.exists():
                print("s%-2d %-11s SKIP  %s missing" % (idx, name, fn))
                skipped += 1
                continue
            ow, oh = blob_size(pic)
            nw, nh = png_size(src)
            if ow and nh and abs((nw / nh) / (ow / oh) - 1) > ASPECT_TOL:
                print("s%-2d %-11s SKIP  aspect %.3f -> %.3f would distort the "
                      "crop" % (idx, name, ow / oh, nw / nh))
                skipped += 1
                continue
            if not a.dry_run:
                _, rId = slide.part.get_or_add_image_part(str(src))
                pic._element.blipFill.blip.set(qn("r:embed"), rId)
            print("s%-2d %-11s <-    %s" % (idx, name, fn))
            done += 1
        for name, why in STALE.get(idx, {}).items():
            pic = find(slide, name)
            if pic is None:
                print("s%-2d %-11s SKIP  shape not found (stale marker)"
                      % (idx, name))
                continue
            if not a.dry_run:
                flag(slide, pic, why, run)
            print("s%-2d %-11s FLAG  %s, no generator" % (idx, name, why))
            flagged += 1

    print("\nrun      : %s" % run)
    print("updated  : %d    flagged: %d    skipped: %d" % (done, flagged, skipped))
    if a.dry_run:
        print("dry run, nothing written")
        return
    if out == deck:
        bak = deck.with_name("%s.bak_%s%s" % (
            deck.stem, datetime.now().strftime("%Y%m%d_%H%M"), deck.suffix))
        shutil.copy2(deck, bak)
        print("backup   : %s" % bak.name)
    prs.save(str(out))
    print("saved    : %s" % out)


if __name__ == "__main__":
    main()
